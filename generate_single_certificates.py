import pandas as pd
from docx import Document

# Load Excel data (skip first row with school name, use second row as header)
df = pd.read_excel("student records.xlsx", header=1)

# Ask for Reg.No from user
reg_input = input("Enter Reg. No: ").strip()

# Filter student record (convert Reg. No to int first to remove decimal, then to string)
student = df[df["Reg. No"].fillna(0).astype(int).astype(str) == reg_input]

if student.empty:
    print("No student found with Reg. No:", reg_input)
    exit()

# Get first (and only) row
row = student.iloc[0]

name = row["Name"]
father = row["Father"]
class_ = row["Class"]
session = row["Session"]
dob = row["DOB"]
gender = str(row["Gender"]).strip().upper()

# Gender logic
if gender == "M":
    relation = "S/O"
    pronoun_subject = "He"
    pronoun_possessive = "His"
else:
    relation = "D/O"
    pronoun_subject = "She"
    pronoun_possessive = "Her"

# Load Word template
doc = Document("bonafide_certificate.docx")

# Replace placeholders
for para in doc.paragraphs:
    text = para.text
    text = text.replace("{{REGNO}}", reg_input)
    text = text.replace("{{NAME}}", name)
    text = text.replace("{{RELATION}}", relation)
    text = text.replace("{{FATHER}}", father)
    text = text.replace("{{CLASS}}", str(class_))
    text = text.replace("{{SESSION}}", session)
    text = text.replace("{{DOB}}", str(dob))
    text = text.replace("{{PRONOUN_SUBJECT}}", pronoun_subject)
    text = text.replace("{{PRONOUN_POSSESSIVE}}", pronoun_possessive)
    para.text = text

# Save certificate
output_filename = f"certificate_{reg_input}.docx"
doc.save(output_filename)

print(f"Certificate generated: {output_filename}")
